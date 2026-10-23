"""Fixed research templates rendered with existing python-docx/LibreOffice.

No shell, URL download, remote templates or macro execution. Each conversion
gets its own temporary directory and LibreOffice profile; memory/concurrency
are bounded independently from document parsing.
"""
from __future__ import annotations

import io
import os
import re
import shutil
import signal
import subprocess
import tempfile
import threading
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

_EXPORT_GATE = threading.BoundedSemaphore(1)
_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
_IMAGE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")


def _run_converter(command: list[str], timeout: int):
    # soffice may fork a child. Kill its entire isolated process group on a
    # timeout, otherwise a timed-out export could outlive the concurrency gate.
    process = subprocess.Popen(
        command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, start_new_session=True,
    )
    try:
        stdout, stderr = process.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        process.communicate()
        raise
    return subprocess.CompletedProcess(command, process.returncode, stdout, stderr)


def _inline(paragraph, text: str):
    for token in _INLINE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            paragraph.add_run(token[1:-1]).font.name = "DejaVu Sans Mono"
        else:
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                # Store the literal target as evidence even for native
                # knowledge:// links, which are not portable browser URLs.
                paragraph.add_run(f"{link[1]} ({link[2]})")
            else:
                paragraph.add_run(token)


def render_docx(title: str, markdown: str, images: dict[str, bytes] | None = None) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "WenQuanYi Zen Hei"
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "WenQuanYi Zen Hei")
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "WenQuanYi Zen Hei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "WenQuanYi Zen Hei")
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    document.core_properties.title = title
    document.core_properties.subject = "Reviewed research"
    document.add_heading(title, 0)
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    footer = section.footer.paragraphs[0]
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = markdown.replace("\r\n", "\n").splitlines()
    # The fixed document title replaces the first Markdown H1 instead of
    # producing adjacent duplicate document titles.
    for index, line in enumerate(lines):
        if not line.strip():
            continue
        if line.startswith("# "):
            lines.pop(index)
        break
    images = images or {}
    i = 0
    in_code = False
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            paragraph = document.add_paragraph()
            paragraph.add_run(line).font.name = "DejaVu Sans Mono"
            i += 1
            continue
        # The fixed template already separates sections with heading spacing.
        # A Markdown horizontal rule should never become visible punctuation.
        if re.fullmatch(r"\s*(?:-{3,}|\*{3,}|_{3,})\s*", line):
            i += 1
            continue
        if i + 1 < len(lines) and "|" in line and _TABLE_RULE.match(lines[i + 1]):
            rows = [line]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(lines[i])
                i += 1
            cells = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
            width = max(len(row) for row in cells)
            if width > 20 or len(cells) > 2000:
                raise ValueError("table exceeds supported research template dimensions")
            table = document.add_table(rows=len(cells), cols=width)
            table.style = "Table Grid"
            for row_num, row in enumerate(cells):
                for col_num, value in enumerate(row):
                    cell = table.cell(row_num, col_num)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    properties = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    margins = OxmlElement("w:tcMar")
                    for edge in ("top", "left", "bottom", "right"):
                        border = OxmlElement(f"w:{edge}")
                        for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                            border.set(qn(f"w:{key}"), val)
                        borders.append(border)
                        margin = OxmlElement(f"w:{edge}")
                        margin.set(qn("w:w"), "100")
                        margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    properties.append(borders)
                    properties.append(margins)
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.15
                    _inline(paragraph, value)
                    if row_num == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "F2F2F2")
                        properties.append(shade)
                        for run in paragraph.runs:
                            run.bold = True
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", line)
        if heading:
            _inline(document.add_heading("", min(3, len(heading[1]))), heading[2])
        elif _IMAGE.search(line):
            cursor = 0
            for match in _IMAGE.finditer(line):
                if match.start() > cursor:
                    _inline(document.add_paragraph(), line[cursor:match.start()])
                data = images.get(match[2])
                if data is None:
                    # Failing is deliberate: never silently export a missing
                    # chart or fetch an attacker-controlled URL from the server.
                    raise ValueError(f"image must be an authorized stored attachment: {match[1]}")
                document.add_picture(io.BytesIO(data), width=Inches(6))
                if match[1]:
                    document.add_paragraph(match[1], style="Caption")
                cursor = match.end()
            if cursor < len(line):
                _inline(document.add_paragraph(), line[cursor:])
        elif line.strip():
            bullet = re.match(r"^\s*[-*+]\s+(.*)", line)
            number = re.match(r"^\s*\d+\.\s+(.*)", line)
            if bullet:
                _inline(document.add_paragraph(style="List Bullet"), bullet[1])
            elif number:
                _inline(document.add_paragraph(style="List Number"), number[1])
            else:
                _inline(document.add_paragraph(), line)
        i += 1
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()


def export_research(title: str, markdown: str, format: str,
                    images: dict[str, bytes] | None = None, timeout: int = 90) -> tuple[bytes, str]:
    if format not in {"md", "docx", "pdf"}:
        raise ValueError("supported formats: md, docx, pdf")
    if len(markdown.encode("utf-8")) > 1024 * 1024:
        raise ValueError("research document exceeds 1 MiB")
    if sum(map(len, (images or {}).values())) > 20 * 1024 * 1024:
        raise ValueError("research images exceed 20 MiB")
    if not _EXPORT_GATE.acquire(timeout=5):
        raise RuntimeError("export converter is busy; retry later")
    try:
        if format == "md":
            return markdown.encode("utf-8"), "text/markdown; charset=utf-8"
        data = render_docx(title, markdown, images)
        if format == "docx":
            return data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        executable = shutil.which("libreoffice") or shutil.which("soffice")
        if not executable:
            raise RuntimeError("LibreOffice is unavailable in the document service")
        with tempfile.TemporaryDirectory(prefix="research-export-") as folder:
            root = Path(folder)
            source = root / "research.docx"
            source.write_bytes(data)
            profile = (root / "profile").as_uri()
            completed = _run_converter(
                [executable, f"-env:UserInstallation={profile}", "--headless",
                 "--norestore", "--nodefault", "--nolockcheck",
                 "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(root), str(source)],
                timeout=timeout,
            )
            output = root / "research.pdf"
            if completed.returncode != 0 or not output.is_file():
                raise RuntimeError("LibreOffice conversion failed")
            result = output.read_bytes()
            if not result.startswith(b"%PDF-"):
                raise RuntimeError("LibreOffice returned an invalid PDF")
            return result, "application/pdf"
    finally:
        _EXPORT_GATE.release()
