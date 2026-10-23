import io
import subprocess
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from PIL import Image

from docreader.research_export import export_research, render_docx


def test_docx_preserves_chinese_table_image_and_citation():
    image = io.BytesIO()
    Image.new("RGB", (20, 10), color="blue").save(image, format="PNG")
    markdown = (
        "## 技术证据\n**中文投资研究**\n\n"
        "| 公司 | 判断 |\n| --- | --- |\n| 甲公司 | 待验证 |\n\n"
        "![实验结果](resource://trusted-image)\n"
        "[原文](knowledge://source/chunks/chunk)\n"
    )
    data, mime = export_research("AI4S 研究", markdown, "docx", {"resource://trusted-image": image.getvalue()})
    assert "wordprocessingml" in mime
    doc = Document(io.BytesIO(data))
    assert "中文投资研究" in "\n".join(p.text for p in doc.paragraphs)
    assert doc.tables[0].cell(1, 0).text == "甲公司"
    assert len(doc.inline_shapes) == 1
    assert "knowledge://source/chunks/chunk" in "\n".join(p.text for p in doc.paragraphs)


def test_export_missing_image_fails_without_network_fetch():
    with pytest.raises(ValueError, match="authorized stored attachment"):
        render_docx("title", "![secret](http://169.254.169.254/latest/meta-data)")


def test_fixed_template_omits_rules_and_duplicate_title():
    data = render_docx(
        "中文研究报告",
        "# 中文研究报告\n\n## 结论\n证据待验证。\n\n---\n\n"
        "## 来源与证据版本\n### 1. 原始研究资料\n"
        "资料编号：01234567-890 · 证据版本：aabbccdd-001\n"
        "写作时资料记录时间：2026-10-20 08:00:00 UTC\n",
    )
    doc = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs]
    assert paragraphs.count("中文研究报告") == 1
    assert "---" not in paragraphs
    assert "来源与证据版本" in paragraphs
    assert any("证据版本：aabbccdd-001" in p for p in paragraphs)
    assert not doc.styles["Title"].element.xpath(".//w:pBdr")


def test_conversion_timeout_releases_gate():
    with patch("docreader.research_export.shutil.which", return_value="/usr/bin/libreoffice"), \
         patch("docreader.research_export._run_converter", side_effect=subprocess.TimeoutExpired("lo", 90)):
        with pytest.raises(subprocess.TimeoutExpired):
            export_research("title", "text", "pdf")
    data, _ = export_research("title", "still available", "docx")
    assert data.startswith(b"PK")


def test_conversions_use_isolated_profiles_and_clean_up():
    profiles = []

    def converter(command, **kwargs):
        profiles.append(next(x for x in command if x.startswith("-env:UserInstallation=")))
        output = Path(command[command.index("--outdir") + 1]) / "research.pdf"
        output.write_bytes(b"%PDF-1.4\n")
        return subprocess.CompletedProcess(command, 0)

    with patch("docreader.research_export.shutil.which", return_value="/usr/bin/libreoffice"), \
         patch("docreader.research_export._run_converter", side_effect=converter):
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: export_research("title", "中文", "pdf"), range(2)))
    assert all(data.startswith(b"%PDF-") for data, _ in results)
    assert len(set(profiles)) == 2
    for profile in profiles:
        assert not Path(profile.split("file://", 1)[1]).parent.exists()


def test_export_rejects_unsupported_format_and_oversize():
    with pytest.raises(ValueError):
        export_research("title", "text", "exe")
    with pytest.raises(ValueError, match="1 MiB"):
        export_research("title", "x" * (1024 * 1024 + 1), "md")


def test_real_libreoffice_pdf_when_available():
    import shutil
    if not (shutil.which("libreoffice") or shutil.which("soffice")):
        pytest.skip("LibreOffice is supplied by the production document image")
    data, mime = export_research("中文研究报告", "## 研究结论\n具备来源的技术证据。\n", "pdf")
    assert data.startswith(b"%PDF-") and len(data) > 1000
    assert mime == "application/pdf"


def test_converter_timeout_terminates_isolated_process():
    import sys
    from docreader.research_export import _run_converter
    with pytest.raises(subprocess.TimeoutExpired):
        _run_converter([sys.executable, "-c", "import time; time.sleep(60)"], timeout=0.05)
